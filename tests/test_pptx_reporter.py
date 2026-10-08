"""Tests for the PowerPoint deck, Confluence exporter and Plotly charts."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation

from supportpulse.analytics.cycle_times import CycleTimeCalculator
from supportpulse.analytics.metrics_engine import MetricsEngine
from supportpulse.analytics.sla_forecaster import SLAForecaster
from supportpulse.dashboard import charts
from supportpulse.reporters.confluence_exporter import ConfluenceExporter
from supportpulse.reporters.pptx_reporter import PptxReporter


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------
class TestPptxReporter:
    def test_build_creates_deck(self, policy, tickets, tmp_path: Path):
        reporter = PptxReporter(policy, output_dir=tmp_path)
        path = reporter.build(tickets)
        assert path.exists()
        assert path.name == "executive_status_deck.pptx"

        presentation = Presentation(path)
        assert len(presentation.slides) == 5

    def test_deck_contains_tables(self, policy, tickets, tmp_path: Path):
        reporter = PptxReporter(policy, output_dir=tmp_path)
        path = reporter.build(tickets)
        presentation = Presentation(path)
        tables = [
            shape
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_table
        ]
        assert len(tables) >= 3

    def test_build_with_explicit_metrics(self, policy, tickets, tmp_path: Path):
        metrics = MetricsEngine(policy).compute(tickets)
        reporter = PptxReporter(policy, output_dir=tmp_path)
        path = reporter.build(tickets, metrics=metrics)
        assert path.exists()

    def test_build_with_no_at_risk_tickets(self, policy, sample_ticket, tmp_path: Path):
        reporter = PptxReporter(policy, output_dir=tmp_path)
        path = reporter.build([sample_ticket])
        presentation = Presentation(path)
        assert len(presentation.slides) == 5

    def test_build_custom_output_path(self, policy, tickets, tmp_path: Path):
        reporter = PptxReporter(policy, output_dir=tmp_path)
        target = tmp_path / "deck.pptx"
        path = reporter.build(tickets, output_path=target)
        assert path == target

    def test_default_output_dir_used(self, policy, tickets, monkeypatch, tmp_path):
        monkeypatch.setenv("SUPPORTPULSE_OUTPUT_DIR", str(tmp_path))
        reporter = PptxReporter(policy)
        path = reporter.build(tickets)
        assert path.parent == tmp_path


# ---------------------------------------------------------------------------
# Confluence
# ---------------------------------------------------------------------------
class TestConfluenceExporter:
    def test_storage_format(self, policy, tickets):
        exporter = ConfluenceExporter(policy, space_key="SUPPORT")
        storage = exporter.to_storage_format(tickets)
        assert "<table>" in storage
        assert 'ac:structured-macro ac:name="info"' in storage
        assert "Executive KPIs" in storage

    def test_storage_format_empty(self, policy):
        exporter = ConfluenceExporter(policy)
        storage = exporter.to_storage_format([])
        assert "No active tickets" in storage

    def test_markdown(self, policy, tickets):
        exporter = ConfluenceExporter(policy)
        markdown = exporter.to_markdown(tickets)
        assert markdown.startswith("# SupportPulse")
        assert "| KPI | Value |" in markdown
        assert "SLA Breach Risk Radar" in markdown

    def test_markdown_empty(self, policy):
        exporter = ConfluenceExporter(policy)
        markdown = exporter.to_markdown([])
        assert "_No active tickets in scope._" in markdown

    def test_export_writes_files(self, policy, tickets, tmp_path: Path):
        exporter = ConfluenceExporter(policy, output_dir=tmp_path)
        paths = exporter.export(tickets)
        assert set(paths) == {"confluence", "markdown"}
        for path in paths.values():
            assert path.exists()
            assert path.read_text(encoding="utf-8")

    def test_export_with_explicit_metrics(self, policy, tickets, tmp_path: Path):
        metrics = MetricsEngine(policy).compute(tickets)
        exporter = ConfluenceExporter(policy, output_dir=tmp_path)
        paths = exporter.export(tickets, metrics=metrics, output_dir=tmp_path / "sub")
        assert paths["markdown"].parent == tmp_path / "sub"

    def test_status_cell_unknown(self):
        assert ConfluenceExporter._status_cell("Unknown") == "<td>Unknown</td>"


# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------
class TestCharts:
    @pytest.fixture()
    def datasets(self, policy, tickets):
        engine = MetricsEngine(policy)
        cycle = CycleTimeCalculator()
        forecaster = SLAForecaster(policy).fit(tickets)
        return {
            "metrics": engine.compute(tickets),
            "cfd": cycle.cumulative_flow(tickets),
            "aging": cycle.aging_wip(tickets),
            "flow": cycle.inflow_outflow(tickets, 14),
            "radar": forecaster.breach_radar(tickets),
            "hotspots": engine.component_hotspots(tickets),
            "ecdf": [t.resolution_hours() for t in tickets if t.resolution_hours()],
        }

    def test_cumulative_flow_figure(self, datasets):
        figure = charts.cumulative_flow_figure(datasets["cfd"])
        assert figure.data

    def test_wip_trend_figure(self, datasets):
        assert charts.wip_trend_figure(datasets["cfd"]).data

    def test_aging_wip_figure(self, datasets):
        assert charts.aging_wip_figure(datasets["aging"]).data

    def test_sla_radar_figure(self, datasets):
        assert charts.sla_radar_figure(datasets["radar"]).data

    def test_inflow_outflow_figure(self, datasets):
        assert charts.inflow_outflow_figure(datasets["flow"]).data

    def test_priority_mix_figure(self, datasets):
        assert charts.priority_mix_figure(datasets["metrics"]).data

    def test_component_bar_figure(self, datasets):
        assert charts.component_bar_figure(datasets["hotspots"]).data

    def test_ecdf_figure(self, datasets):
        assert charts.ecdf_figure(datasets["ecdf"]).data

    def test_empty_inputs_render_placeholder(self):
        assert charts.cumulative_flow_figure(pd.DataFrame()).data == ()
        assert charts.wip_trend_figure(pd.DataFrame()).data == ()
        assert charts.aging_wip_figure(pd.DataFrame()).data == ()
        assert charts.sla_radar_figure(pd.DataFrame()).data == ()
        assert charts.inflow_outflow_figure(pd.DataFrame()).data == ()
        assert charts.ecdf_figure([]).data == ()
