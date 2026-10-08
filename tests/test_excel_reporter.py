"""Tests for the openpyxl Excel workbook generator."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from supportpulse.reporters.excel_reporter import ExcelReporter, _is_breach, metrics_rows
from supportpulse.models.metrics import SupportMetrics
from supportpulse.models.sla import SLAPolicy


class TestExcelReporter:
    def test_build_creates_workbook(self, policy, tickets, tmp_path: Path):
        reporter = ExcelReporter(policy, output_dir=tmp_path)
        path = reporter.build(tickets)
        assert path.exists()
        assert path.name == "weekly_support_status.xlsx"

        workbook = load_workbook(path)
        assert set(workbook.sheetnames) == {
            "Executive KPI",
            "Ticket Ledger",
            "Pivot Analysis",
            "SLA Heatmap",
        }

    def test_executive_sheet_content(self, policy, tickets, tmp_path: Path):
        reporter = ExcelReporter(policy, output_dir=tmp_path)
        path = reporter.build(tickets)
        sheet = load_workbook(path)["Executive KPI"]
        values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
        assert any("SupportPulse" in str(value) for value in values)
        assert any("Total tickets" == str(value) for value in values)

    def test_heatmap_has_conditional_formatting(self, policy, tickets, tmp_path: Path):
        reporter = ExcelReporter(policy, output_dir=tmp_path)
        path = reporter.build(tickets)
        sheet = load_workbook(path)["SLA Heatmap"]
        rules = list(sheet.conditional_formatting)
        assert rules, "Expected colour-scale rules on the SLA heatmap"

    def test_ledger_frozen_header(self, policy, tickets, tmp_path: Path):
        reporter = ExcelReporter(policy, output_dir=tmp_path)
        path = reporter.build(tickets)
        sheet = load_workbook(path)["Ticket Ledger"]
        assert sheet.freeze_panes == "A2"

    def test_build_with_explicit_metrics(self, policy, tickets, tmp_path: Path):
        from supportpulse.analytics.metrics_engine import MetricsEngine

        metrics = MetricsEngine(policy).compute(tickets)
        reporter = ExcelReporter(policy, output_dir=tmp_path)
        path = reporter.build(tickets, metrics=metrics)
        assert path.exists()

    def test_build_with_custom_output_path(self, policy, tickets, tmp_path: Path):
        reporter = ExcelReporter(policy, output_dir=tmp_path)
        target = tmp_path / "nested" / "custom.xlsx"
        path = reporter.build(tickets, output_path=target)
        assert path == target
        assert target.exists()

    def test_build_empty_population(self, policy, tmp_path: Path):
        reporter = ExcelReporter(policy, output_dir=tmp_path)
        path = reporter.build([])
        workbook = load_workbook(path)
        assert "Executive KPI" in workbook.sheetnames
        ledger = workbook["Ticket Ledger"]
        assert ledger["A1"].value == "No tickets in scope"

    def test_rag_rules_on_compliance(self, policy, tickets, tmp_path: Path):
        reporter = ExcelReporter(policy, output_dir=tmp_path)
        path = reporter.build(tickets)
        sheet = load_workbook(path)["Executive KPI"]
        assert list(sheet.conditional_formatting)


class TestExcelHelpers:
    def test_metrics_rows(self):
        metrics = SupportMetrics(
            by_priority={
                "P2-High": {
                    "resolved": 10,
                    "mttr_hours": 12.5,
                    "p85_hours": 20.0,
                    "breaches": 2,
                    "breach_rate": 0.2,
                    "sla_target_hours": 24.0,
                }
            }
        )
        rows = metrics_rows(metrics)
        assert rows["P2-High"]["mttr"] == 12.5
        assert rows["P2-High"]["target"] == 24.0

    def test_is_breach(self, policy):
        within = pd.Series({"priority": "P2-High", "resolution_hours": 10.0})
        over = pd.Series({"priority": "P2-High", "resolution_hours": 30.0})
        missing = pd.Series({"priority": "P2-High", "resolution_hours": None})
        assert _is_breach(within, policy) == 0.0
        assert _is_breach(over, policy) == 1.0
        assert _is_breach(missing, policy) == 0.0

    def test_default_output_dir_used(self, policy, tickets, monkeypatch, tmp_path):
        monkeypatch.setenv("SUPPORTPULSE_OUTPUT_DIR", str(tmp_path))
        reporter = ExcelReporter(policy)
        path = reporter.build(tickets)
        assert path.parent == tmp_path
