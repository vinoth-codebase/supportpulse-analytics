"""Executive PowerPoint status deck generation with python-pptx."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from supportpulse.analytics.cycle_times import CycleTimeCalculator
from supportpulse.analytics.metrics_engine import MetricsEngine
from supportpulse.analytics.sla_forecaster import SLAForecaster
from supportpulse.config import resolve_output_dir
from supportpulse.models.metrics import SupportMetrics
from supportpulse.models.sla import SLAPolicy
from supportpulse.models.ticket import JiraTicket

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xD9, 0xE2, 0xF3)
GREY = RGBColor(0xF2, 0xF2, 0xF2)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xED, 0x7D, 0x31)
RED = RGBColor(0xC0, 0x39, 0x2B)

RAG_COLORS = {"Green": GREEN, "Amber": AMBER, "Red": RED}


class PptxReporter:
    """Builds the ``executive_status_deck.pptx`` visual-management deck."""

    def __init__(
        self,
        sla_policy: SLAPolicy | None = None,
        output_dir: str | Path | None = None,
        hours_per_day: float = 8.0,
    ) -> None:
        self.policy = sla_policy or SLAPolicy.default()
        self.engine = MetricsEngine(self.policy, hours_per_day=hours_per_day)
        self.cycle_times = CycleTimeCalculator(hours_per_day)
        self.forecaster = SLAForecaster(self.policy, hours_per_day=hours_per_day)
        self.output_dir = Path(output_dir) if output_dir else resolve_output_dir()

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------
    def build(
        self,
        tickets: Iterable[JiraTicket],
        metrics: SupportMetrics | None = None,
        output_path: str | Path | None = None,
    ) -> Path:
        tickets = list(tickets)
        metrics = metrics or self.engine.compute(tickets)
        self.forecaster.fit(tickets)

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

        self._title_slide(presentation, metrics)
        self._scorecard_slide(presentation, metrics)
        self._flow_slide(presentation, tickets)
        self._components_slide(presentation, tickets)
        self._risk_slide(presentation, tickets)

        destination = (
            Path(output_path)
            if output_path
            else self.output_dir / "executive_status_deck.pptx"
        )
        destination.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(destination)
        return destination

    # ------------------------------------------------------------------
    # Slide builders
    # ------------------------------------------------------------------
    def _title_slide(self, presentation: Presentation, metrics: SupportMetrics) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._band(slide, "SupportPulse-Analytics", "Visual Management — Executive Status Deck")

        box = slide.shapes.add_textbox(Inches(1.0), Inches(3.1), Inches(11.3), Inches(2.2))
        frame = box.text_frame
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.text = "Weekly Support Operations Review"
        paragraph.font.size = Pt(30)
        paragraph.font.bold = True
        paragraph.font.color.rgb = NAVY

        sub = frame.add_paragraph()
        sub.text = (
            f"Reporting window: {metrics.window_days} days   |   "
            f"{metrics.total_tickets} tickets   |   "
            f"MTTR {metrics.mttr_hours:.1f} business hrs   |   "
            f"SLA breach rate {metrics.breach_rate * 100:.1f}%"
        )
        sub.font.size = Pt(15)
        sub.font.color.rgb = BLUE

        stamp = frame.add_paragraph()
        stamp.text = f"Generated {metrics.generated_at.strftime('%d %b %Y %H:%M UTC')}"
        stamp.font.size = Pt(12)
        stamp.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    def _scorecard_slide(self, presentation: Presentation, metrics: SupportMetrics) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._header(slide, "Weekly Support Executive Scorecard")

        rows = metrics.as_summary_rows()
        table = self._table(slide, rows=len(rows) + 1, cols=2, top=1.6, width=6.0, left=0.8)
        self._set_header(table, ["KPI", "Value"])
        for index, (label, value) in enumerate(rows, start=1):
            table.cell(index, 0).text = label
            table.cell(index, 1).text = value
            self._shade(table.cell(index, 0), GREY if index % 2 else WHITE)

        compliance = self.engine.sla_compliance_table_from_metrics(metrics)
        side = self._table(
            slide, rows=len(compliance) + 1, cols=3, top=1.6, width=5.5, left=7.2
        )
        self._set_header(side, ["Priority", "Attainment", "Status"])
        for index, (_, row) in enumerate(compliance.iterrows(), start=1):
            side.cell(index, 0).text = str(row["priority"])
            side.cell(index, 1).text = f"{row['attainment'] * 100:.1f}%"
            cell = side.cell(index, 2)
            cell.text = str(row["status"])
            colour = RAG_COLORS.get(str(row["status"]))
            if colour is not None:
                self._fill(cell, colour)

    def _flow_slide(self, presentation: Presentation, tickets: list[JiraTicket]) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._header(slide, "Ticket Inflow vs. Outflow Trends")

        flow = self.cycle_times.inflow_outflow(tickets, window_days=14)
        chart_data = CategoryChartData()
        chart_data.categories = [stamp.strftime("%d %b") for stamp in flow.index]
        chart_data.add_series("Inflow", tuple(flow["inflow"].tolist()))
        chart_data.add_series("Outflow", tuple(flow["outflow"].tolist()))

        graphic_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(0.8),
            Inches(1.5),
            Inches(11.7),
            Inches(5.4),
            chart_data,
        )
        chart = graphic_frame.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    def _components_slide(self, presentation: Presentation, tickets: list[JiraTicket]) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._header(slide, "Top Recurring Defect Components")

        hotspots = self.engine.component_hotspots(tickets, top_n=8)
        table = self._table(
            slide, rows=len(hotspots) + 1, cols=5, top=1.6, width=11.7, left=0.8
        )
        self._set_header(
            table, ["Component", "Tickets", "Resolved", "MTTR (hrs)", "Breach rate"]
        )
        for index, (_, row) in enumerate(hotspots.iterrows(), start=1):
            table.cell(index, 0).text = str(row["component"])
            table.cell(index, 1).text = str(int(row["tickets"]))
            table.cell(index, 2).text = str(int(row["resolved"]))
            table.cell(index, 3).text = f"{row['mttr_hours']:.1f}"
            table.cell(index, 4).text = f"{row['breach_rate'] * 100:.1f}%"
            if index % 2 == 0:
                for col in range(5):
                    self._shade(table.cell(index, col), GREY)

    def _risk_slide(self, presentation: Presentation, tickets: list[JiraTicket]) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._header(slide, "Critical At-Risk Backlog Items")

        radar = self.forecaster.breach_radar(tickets)
        critical = radar[radar["risk_level"].isin(["Critical", "High"])].head(10)
        if critical.empty:
            box = slide.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(11), Inches(1.2))
            box.text_frame.text = "No tickets currently flagged at high or critical breach risk."
            return

        columns = ["key", "priority", "status", "runway_hours", "breach_probability", "risk_level"]
        table = self._table(
            slide, rows=len(critical) + 1, cols=len(columns), top=1.6, width=11.7, left=0.8
        )
        self._set_header(
            table,
            ["Key", "Priority", "Status", "Runway (hrs)", "Breach %", "Risk"],
        )
        for index, (_, row) in enumerate(critical.iterrows(), start=1):
            table.cell(index, 0).text = str(row["key"])
            table.cell(index, 1).text = str(row["priority"])
            table.cell(index, 2).text = str(row["status"])
            table.cell(index, 3).text = f"{row['runway_hours']:.1f}"
            table.cell(index, 4).text = f"{row['breach_probability']:.1f}%"
            cell = table.cell(index, 5)
            cell.text = str(row["risk_level"])
            colour = RAG_COLORS.get("Red" if row["risk_level"] == "Critical" else "Amber")
            if colour is not None:
                self._fill(cell, colour)

    # ------------------------------------------------------------------
    # Styling helpers
    # ------------------------------------------------------------------
    def _band(self, slide, title: str, subtitle: str) -> None:
        shape = slide.shapes.add_shape(
            1, Emu(0), Emu(0), Inches(13.333), Inches(1.4)  # MSO_SHAPE.RECTANGLE
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY
        shape.line.fill.background()
        frame = shape.text_frame
        frame.word_wrap = True
        frame.text = title
        frame.paragraphs[0].font.size = Pt(30)
        frame.paragraphs[0].font.bold = True
        frame.paragraphs[0].font.color.rgb = WHITE
        para = frame.add_paragraph()
        para.text = subtitle
        para.font.size = Pt(14)
        para.font.color.rgb = LIGHT

    def _header(self, slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.9))
        frame = box.text_frame
        frame.text = text
        frame.paragraphs[0].font.size = Pt(26)
        frame.paragraphs[0].font.bold = True
        frame.paragraphs[0].font.color.rgb = NAVY
        frame.paragraphs[0].alignment = PP_ALIGN.LEFT

    def _table(self, slide, rows: int, cols: int, top: float, width: float, left: float):
        shape = slide.shapes.add_table(
            rows, cols, Inches(left), Inches(top), Inches(width), Inches(0.35 * rows)
        )
        return shape.table

    def _set_header(self, table, labels: list[str]) -> None:
        for index, label in enumerate(labels):
            cell = table.cell(0, index)
            cell.text = label
            self._fill(cell, NAVY)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.bold = True
                paragraph.font.size = Pt(13)
                paragraph.font.color.rgb = WHITE

    @staticmethod
    def _fill(cell, colour: RGBColor) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = colour

    @staticmethod
    def _shade(cell, colour: RGBColor) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = colour
