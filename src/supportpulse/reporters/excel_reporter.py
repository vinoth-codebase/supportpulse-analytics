"""Formatted multi-tab Excel workbook generation with openpyxl."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from supportpulse.analytics.metrics_engine import MetricsEngine
from supportpulse.analytics.sla_forecaster import SLAForecaster
from supportpulse.config import resolve_output_dir
from supportpulse.models.metrics import SupportMetrics
from supportpulse.models.sla import SLAPolicy
from supportpulse.models.ticket import JiraTicket

# Palette -------------------------------------------------------------------
NAVY = "1F3864"
BLUE = "2E75B6"
LIGHT = "D9E2F3"
GREEN = "C6EFCE"
AMBER = "FFEB9C"
RED = "FFC7CE"
GREY = "F2F2F2"

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class ExcelReporter:
    """Builds ``weekly_support_status.xlsx`` with KPI, ledger and heatmaps."""

    def __init__(
        self,
        sla_policy: SLAPolicy | None = None,
        output_dir: str | Path | None = None,
        hours_per_day: float = 8.0,
    ) -> None:
        self.policy = sla_policy or SLAPolicy.default()
        self.engine = MetricsEngine(self.policy, hours_per_day=hours_per_day)
        self.forecaster = SLAForecaster(self.policy, hours_per_day=hours_per_day)
        self.output_dir = Path(output_dir) if output_dir else resolve_output_dir()

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------
    def build(
        self,
        tickets: Iterable[JiraTicket],
        metrics: SupportMetrics | None = None,
        output_path: str | Path | None = None,
    ) -> Path:
        """Generate the workbook and return the written path."""

        tickets = list(tickets)
        metrics = metrics or self.engine.compute(tickets)
        self.forecaster.fit(tickets)

        workbook = Workbook()
        workbook.remove(workbook.active)

        self._executive_sheet(workbook, tickets, metrics)
        self._ledger_sheet(workbook, tickets)
        self._pivot_sheet(workbook, tickets)
        self._heatmap_sheet(workbook, tickets)

        destination = Path(output_path) if output_path else self.output_dir / "weekly_support_status.xlsx"
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
        return destination

    # ------------------------------------------------------------------
    # Sheet builders
    # ------------------------------------------------------------------
    def _executive_sheet(
        self, workbook: Workbook, tickets: list[JiraTicket], metrics: SupportMetrics
    ) -> None:
        sheet = workbook.create_sheet("Executive KPI")
        sheet.sheet_view.showGridLines = False
        self._title(sheet, "SupportPulse — Weekly Support Status", span=4)

        sheet["A3"] = "Generated"
        sheet["B3"] = metrics.generated_at.strftime("%d %b %Y %H:%M UTC")
        sheet["A4"] = "Reporting window"
        sheet["B4"] = f"{metrics.window_days} days"
        for row in (3, 4):
            sheet.cell(row=row, column=1).font = Font(bold=True, color=NAVY)

        # KPI card -----------------------------------------------------
        start = 6
        self._section(sheet, "Headline KPIs", start, span=4)
        row = start + 1
        for label, value in metrics.as_summary_rows():
            sheet.cell(row=row, column=1, value=label).font = Font(bold=False)
            cell = sheet.cell(row=row, column=2, value=value)
            cell.alignment = Alignment(horizontal="right")
            for column in (1, 2):
                sheet.cell(row=row, column=column).border = BORDER
                if row % 2 == 0:
                    sheet.cell(row=row, column=column).fill = PatternFill(
                        "solid", fgColor=GREY
                    )
            row += 1

        # SLA compliance table -----------------------------------------
        row += 1
        self._section(sheet, "SLA Attainment by Priority", row, span=7)
        row += 1
        compliance = self.engine.sla_compliance_table(tickets)
        row = self._write_frame(sheet, compliance, row)

        # Risk summary -------------------------------------------------
        row += 1
        self._section(sheet, "SLA Breach Risk Radar", row, span=7)
        row += 1
        radar = self.forecaster.breach_radar(tickets)
        if not radar.empty:
            row = self._write_frame(sheet, radar.head(15), row)

        self._autosize(sheet)

    def _ledger_sheet(self, workbook: Workbook, tickets: list[JiraTicket]) -> None:
        sheet = workbook.create_sheet("Ticket Ledger")
        frame = self.engine.to_frame(tickets)
        if frame.empty:
            sheet["A1"] = "No tickets in scope"
            return
        columns = [
            "key",
            "priority",
            "status",
            "component",
            "assignee",
            "created_at",
            "resolved_at",
            "resolution_hours",
            "age_hours",
            "first_response_hours",
            "reassignments",
        ]
        frame = frame[[c for c in columns if c in frame.columns]].copy()
        frame["created_at"] = pd.to_datetime(frame["created_at"]).dt.strftime(
            "%Y-%m-%d %H:%M"
        )
        frame["resolved_at"] = pd.to_datetime(frame["resolved_at"]).dt.strftime(
            "%Y-%m-%d %H:%M"
        )
        self._write_frame(sheet, frame, start_row=1, freeze="A2")
        self._autosize(sheet)

    def _pivot_sheet(self, workbook: Workbook, tickets: list[JiraTicket]) -> None:
        sheet = workbook.create_sheet("Pivot Analysis")
        frame = self.engine.to_frame(tickets)
        if frame.empty:
            sheet["A1"] = "No tickets in scope"
            return

        pivot = pd.crosstab(frame["priority"], frame["status"])
        pivot["Total"] = pivot.sum(axis=1)
        pivot.loc["Total"] = pivot.sum(axis=0)
        sheet["A1"] = "Tickets by Priority x Status"
        sheet["A1"].font = Font(bold=True, size=13, color=NAVY)
        self._write_frame(sheet, pivot.reset_index(), start_row=3)

        start = len(pivot) + 6
        sheet.cell(row=start, column=1, value="Resolution Time by Priority (business hours)")
        sheet.cell(row=start, column=1).font = Font(bold=True, size=13, color=NAVY)
        summary = pd.DataFrame(metrics_rows(self.engine.compute(tickets))).T.reset_index()
        summary.columns = ["priority", "mttr", "p85", "breaches", "breach_rate", "target"]
        self._write_frame(sheet, summary, start_row=start + 1)
        self._autosize(sheet)

    def _heatmap_sheet(self, workbook: Workbook, tickets: list[JiraTicket]) -> None:
        sheet = workbook.create_sheet("SLA Heatmap")
        sheet.sheet_view.showGridLines = False
        self._title(sheet, "SLA Heatmap — Breach Rate by Component / Priority", span=6)

        frame = self.engine.to_frame(tickets)
        if frame.empty:
            sheet["A3"] = "No tickets in scope"
            return

        # Build component x priority breach-rate matrix.
        records = []
        for _, row in frame.iterrows():
            records.append(
                {
                    "component": row["component"] or "Unassigned",
                    "priority": row["priority"],
                    "breach": _is_breach(row, self.policy),
                }
            )
        matrix = pd.DataFrame(records)
        heat = (
            matrix.pivot_table(
                index="component", columns="priority", values="breach", aggfunc="mean"
            )
            .fillna(0.0)
            .round(3)
        )
        heat["Overall"] = heat.mean(axis=1).round(3)
        heat = heat.reset_index()

        start_row = 3
        first_data_row = start_row + 1
        self._write_frame(sheet, heat, start_row=start_row)

        last_row = first_data_row + len(heat) - 1
        first_col = 2
        last_col = first_col + len(heat.columns) - 2  # exclude 'component'
        data_range = (
            f"{get_column_letter(first_col)}{first_data_row}:"
            f"{get_column_letter(last_col)}{last_row}"
        )
        sheet.conditional_formatting.add(
            data_range,
            ColorScaleRule(
                start_type="num",
                start_value=0.0,
                start_color=GREEN,
                mid_type="num",
                mid_value=0.15,
                mid_color=AMBER,
                end_type="num",
                end_value=0.4,
                end_color=RED,
            ),
        )
        for row in sheet[data_range]:
            for cell in row:
                cell.number_format = "0.0%"

        self._autosize(sheet)

    # ------------------------------------------------------------------
    # Styling helpers
    # ------------------------------------------------------------------
    def _title(self, sheet: Worksheet, text: str, span: int = 6) -> None:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
        cell = sheet.cell(row=1, column=1, value=text)
        cell.font = Font(bold=True, size=16, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 26

    def _section(self, sheet: Worksheet, text: str, row: int, span: int = 6) -> None:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
        cell = sheet.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, size=12, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.alignment = Alignment(horizontal="left")

    def _write_frame(
        self,
        sheet: Worksheet,
        frame: pd.DataFrame,
        start_row: int,
        freeze: str | None = None,
    ) -> int:
        """Write a DataFrame starting at *start_row*; return the next free row."""

        if frame.empty:
            return start_row + 1
        header = list(frame.columns)
        for col_index, name in enumerate(header, start=1):
            cell = sheet.cell(row=start_row, column=col_index, value=str(name))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER

        for offset, (_, row) in enumerate(frame.iterrows(), start=1):
            for col_index, name in enumerate(header, start=1):
                value = row[name]
                if isinstance(value, (pd.Timestamp, datetime)):
                    value = value.strftime("%Y-%m-%d %H:%M")
                elif hasattr(value, "item"):
                    value = value.item()
                cell = sheet.cell(row=start_row + offset, column=col_index, value=value)
                cell.border = BORDER
                if offset % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=GREY)

        # RAG conditional formatting for any 'status' column.
        if "status" in header:
            column_letter = get_column_letter(header.index("status") + 1)
            rng = f"{column_letter}{start_row + 1}:{column_letter}{start_row + len(frame)}"
            sheet.conditional_formatting.add(
                rng, CellIsRule(operator="equal", formula=['"Green"'], fill=PatternFill("solid", fgColor=GREEN))
            )
            sheet.conditional_formatting.add(
                rng, CellIsRule(operator="equal", formula=['"Amber"'], fill=PatternFill("solid", fgColor=AMBER))
            )
            sheet.conditional_formatting.add(
                rng, CellIsRule(operator="equal", formula=['"Red"'], fill=PatternFill("solid", fgColor=RED))
            )

        if freeze:
            sheet.freeze_panes = freeze
        return start_row + len(frame) + 2

    def _autosize(self, sheet: Worksheet, min_width: int = 12, max_width: int = 60) -> None:
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            longest = 0
            for cell in column_cells:
                if cell.value is None:
                    continue
                longest = max(longest, len(str(cell.value)))
            sheet.column_dimensions[letter].width = min(max_width, max(min_width, longest + 3))


def metrics_rows(metrics: SupportMetrics) -> dict[str, dict[str, float]]:
    """Flatten ``by_priority`` metrics into a simple mapping."""

    rows: dict[str, dict[str, float]] = {}
    for priority, values in metrics.by_priority.items():
        rows[priority] = {
            "mttr": values.get("mttr_hours", 0.0),
            "p85": values.get("p85_hours", 0.0),
            "breaches": values.get("breaches", 0.0),
            "breach_rate": values.get("breach_rate", 0.0),
            "target": values.get("sla_target_hours", 0.0),
        }
    return rows


def _is_breach(row: pd.Series, policy: SLAPolicy) -> float:
    hours = row.get("resolution_hours")
    if hours is None or pd.isna(hours):
        return 0.0
    target = policy.target_for(row["priority"]).resolution_hours
    return 1.0 if hours > target else 0.0
